"""
회의록 md 파일을 회사 표준 서식(회의록 양식.docx)에 맞춰 docx로 변환한다.

기대하는 md 구조:

---
회의명: ...
일시: ...
장소: ...
참석자: ...
작성자: ...
---
# 1. 주요 논의 안건
- 안건1
- 안건2

# 2. 주요 논의 내용
## 가. 주제1
내용...
## 나. 주제2
내용...

# 3. 결정 사항 및 향후 일정
- 항목명: 내용 설명
"""

import argparse
import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

KOREAN_LABELS = list("가나다라마바사아자차카타파하거너더러머버서어저처커터퍼허")


def parse_md(md_path: Path):
    text = md_path.read_text(encoding="utf-8")

    meta = {}
    body = text
    if text.lstrip().startswith("---"):
        stripped = text.lstrip()
        end = stripped.find("\n---", 3)
        if end != -1:
            front = stripped[3:end].strip()
            body = stripped[end + 4 :]
            for line in front.splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    meta[k.strip()] = v.strip()

    def extract_section(heading_text: str) -> str:
        pattern = rf"^#\s*{re.escape(heading_text)}\s*$(.*?)(?=^#\s|\Z)"
        m = re.search(pattern, body, re.S | re.M)
        return m.group(1).strip() if m else ""

    agenda_raw = extract_section("1. 주요 논의 안건")
    content_raw = extract_section("2. 주요 논의 내용")
    decisions_raw = extract_section("3. 결정 사항 및 향후 일정")

    def bullet_items(raw: str) -> list[str]:
        return [re.sub(r"^[-*]\s*", "", l).strip() for l in raw.splitlines() if l.strip()]

    agenda_items = bullet_items(agenda_raw)
    decision_items = bullet_items(decisions_raw)

    subsections = []
    blocks = re.split(r"\n(?=##\s)", content_raw)
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        m = re.match(r"##\s*(.+)", block)
        if m:
            title = re.sub(r"^[가-힣]\.\s*", "", m.group(1).strip())
            rest = block[m.end() :].strip()
        else:
            title, rest = "", block
        subsections.append((title, rest))

    return meta, agenda_items, subsections, decision_items


def clone_after(anchor: Paragraph, style_source: Paragraph) -> Paragraph:
    new_p = copy.deepcopy(style_source._p)
    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def set_text(paragraph: Paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def fill_list(placeholder: Paragraph, items: list[str]) -> Paragraph:
    if not items:
        set_text(placeholder, "")
        return placeholder
    set_text(placeholder, items[0])
    anchor = placeholder
    for item in items[1:]:
        anchor = clone_after(anchor, placeholder)
        set_text(anchor, item)
    return anchor


def find_paragraph_index(doc: docx.Document, text: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    raise ValueError(f"템플릿에서 문단을 찾지 못했습니다: {text!r}")


def set_table_field(table, label_no_space: str, value: str) -> bool:
    for row in table.rows:
        if row.cells[0].text.replace(" ", "").strip() == label_no_space:
            row.cells[1].text = value
            return True
    return False


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def add_decisions_table(doc: docx.Document, after_paragraph: Paragraph, items: list[str]):
    table = doc.add_table(rows=1, cols=2)
    try:
        table.style = "Table Grid"
    except KeyError:
        set_table_borders(table)
    hdr = table.rows[0].cells
    hdr[0].text = "항목"
    hdr[1].text = "내용"

    for item in items:
        if ":" in item:
            label, value = item.split(":", 1)
        elif "：" in item:
            label, value = item.split("：", 1)
        else:
            label, value = "", item
        row = table.add_row().cells
        row[0].text = label.strip()
        row[1].text = value.strip()

    after_paragraph._p.addnext(table._tbl)


def build_docx(template_path: Path, md_path: Path, output_path: Path):
    meta, agenda_items, subsections, decision_items = parse_md(md_path)

    doc = docx.Document(str(template_path))

    # 부제목(회의명)
    subtitle_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "(회의명)":
            subtitle_idx = i
            break
    if subtitle_idx is not None:
        set_text(doc.paragraphs[subtitle_idx], meta.get("회의명", ""))

    # 정보 표 (회의명/일시/장소/참석자/작성자)
    info_table = doc.tables[0]
    for key in ["회의명", "일시", "장소", "참석자", "작성자"]:
        set_table_field(info_table, key, meta.get(key, ""))

    # 1. 주요 논의 안건
    idx1 = find_paragraph_index(doc, "1. 주요 논의 안건")
    agenda_placeholder = doc.paragraphs[idx1 + 1]
    fill_list(agenda_placeholder, agenda_items)

    # 2. 주요 논의 내용 (가/나/다 ... 소제목 + 내용)
    idx2 = find_paragraph_index(doc, "2. 주요 논의 내용")
    heading2_placeholder = doc.paragraphs[idx2 + 1]
    content_placeholder = doc.paragraphs[idx2 + 2]

    last_anchor = heading2_placeholder
    for n, (title, body_text) in enumerate(subsections):
        label = KOREAN_LABELS[n] if n < len(KOREAN_LABELS) else str(n + 1)
        if n == 0:
            heading_para = heading2_placeholder
        else:
            heading_para = clone_after(last_anchor, heading2_placeholder)
        set_text(heading_para, f"{label}. {title}")
        last_anchor = heading_para

        lines = [l.strip() for l in body_text.splitlines() if l.strip()] or [""]
        for m, line in enumerate(lines):
            if n == 0 and m == 0:
                content_para = content_placeholder
            else:
                content_para = clone_after(last_anchor, content_placeholder)
            set_text(content_para, re.sub(r"^[-*]\s*", "", line))
            last_anchor = content_para

    if not subsections:
        set_text(heading2_placeholder, "")
        set_text(content_placeholder, "")

    # 3. 결정 사항 및 향후 일정 (표 형식 — 기존 완성본 사례를 따름)
    idx3 = find_paragraph_index(doc, "3. 결정 사항 및 향후 일정")
    heading3_para = doc.paragraphs[idx3]
    decisions_placeholder = doc.paragraphs[idx3 + 1]
    # 원본 빈 줄 placeholder는 표로 대체하므로 문단 자체를 제거한다
    decisions_placeholder._p.getparent().remove(decisions_placeholder._p)
    add_decisions_table(doc, heading3_para, decision_items)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def main():
    parser = argparse.ArgumentParser(description="회의록 md를 회사 서식 docx로 변환합니다.")
    parser.add_argument("md_file", help="입력 회의록 md 파일 경로")
    parser.add_argument(
        "--template",
        default="template.docx",
        help="회의록 양식(docx) 템플릿 경로",
    )
    parser.add_argument("--output", help="출력 docx 경로 (기본값: md 파일과 같은 폴더/이름)")
    args = parser.parse_args()

    md_path = Path(args.md_file)
    template_path = Path(args.template)
    if not md_path.is_file():
        print(f"[오류] md 파일을 찾을 수 없습니다: {md_path}", file=sys.stderr)
        sys.exit(1)
    if not template_path.is_file():
        print(f"[오류] 템플릿 파일을 찾을 수 없습니다: {template_path}", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else md_path.with_suffix(".docx")

    build_docx(template_path, md_path, output_path)
    print(f"완료: {output_path}")


if __name__ == "__main__":
    main()
